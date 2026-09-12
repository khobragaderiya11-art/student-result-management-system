"""
Student Result Management System
Python Tkinter (GUI) + Excel (openpyxl) storage

Menu:
  [1] Add Student   -> enter details + marks, auto-calculates Total/Percentage/Result, saves to Excel
  [2] Get Result    -> search a student by Roll No.
  [3] Show All      -> view every saved student in a table
  [4] Exit

Excel file: student_results.xlsx (created automatically next to this script)
Columns: Name | Roll No. | Class | Subject 1 | Subject 2 | Subject 3 | Subject 4 | Subject 5
         | Total Marks | Percentage | Result

Run with:  python student_result_system.py
Requires:  pip install openpyxl
"""

import os
import tkinter as tk
from tkinter import ttk, messagebox

from openpyxl import Workbook, load_workbook

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "student_results.xlsx")
SUBJECT_COUNT = 5
MAX_MARKS_PER_SUBJECT = 100
PASS_PERCENTAGE = 40  # overall percentage needed to Pass

HEADERS = (
    ["Name", "Roll No.", "Class"]
    + [f"Subject {i}" for i in range(1, SUBJECT_COUNT + 1)]
    + ["Total Marks", "Percentage", "Result"]
)


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------
def ensure_excel_file():
    """Create the Excel file with headers if it doesn't already exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(HEADERS)
        wb.save(EXCEL_FILE)


def roll_no_exists(roll_no):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is not None and str(row[1]) == str(roll_no):
            return True
    return False


def append_student_row(row_values):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append(row_values)
    wb.save(EXCEL_FILE)


def find_student_by_roll(roll_no):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is not None and str(row[1]) == str(roll_no):
            return row
    return None


def read_all_students():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))


# ---------------------------------------------------------------------------
# Main Application
# ---------------------------------------------------------------------------
class ResultApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Student Result Management System")
        self.geometry("760x540")
        self.minsize(680, 480)
        self.configure(bg="#eef1f5")

        ensure_excel_file()
        self._build_style()
        self._build_menu_bar()
        self._build_container()

        self.show_add_student()

    # ------------------------------------------------------------------
    def _build_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background="#eef1f5")
        style.configure("TLabel", background="#eef1f5", font=("Segoe UI", 10))
        style.configure("Title.TLabel", background="#eef1f5", font=("Segoe UI", 15, "bold"), foreground="#1c2b4a")
        style.configure("TButton", font=("Segoe UI", 10, "bold"), padding=6)
        style.configure("Menu.TButton", font=("Segoe UI", 10, "bold"), padding=10)
        style.configure("Treeview", rowheight=25, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("Pass.TLabel", background="#eef1f5", foreground="#1f7a3d", font=("Segoe UI", 11, "bold"))
        style.configure("Fail.TLabel", background="#eef1f5", foreground="#b3261e", font=("Segoe UI", 11, "bold"))

    def _build_menu_bar(self):
        bar = ttk.Frame(self, padding=(14, 12))
        bar.pack(fill="x")

        ttk.Label(bar, text="STUDENT RESULT MANAGEMENT", style="Title.TLabel").pack(side="left")

        btns = ttk.Frame(bar)
        btns.pack(side="right")
        ttk.Button(btns, text="1. Add Student", style="Menu.TButton", command=self.show_add_student).pack(side="left", padx=4)
        ttk.Button(btns, text="2. Get Result", style="Menu.TButton", command=self.show_get_result).pack(side="left", padx=4)
        ttk.Button(btns, text="3. Show All Results", style="Menu.TButton", command=self.show_all_results).pack(side="left", padx=4)
        ttk.Button(btns, text="4. Exit", style="Menu.TButton", command=self.destroy).pack(side="left", padx=4)

        ttk.Separator(self, orient="horizontal").pack(fill="x")

    def _build_container(self):
        self.container = ttk.Frame(self, padding=16)
        self.container.pack(fill="both", expand=True)

    def _clear_container(self):
        for widget in self.container.winfo_children():
            widget.destroy()

    # ==================================================================
    # 1. ADD STUDENT
    # ==================================================================
    def show_add_student(self):
        self._clear_container()

        ttk.Label(self.container, text="Add Student", style="Title.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 14))

        fields = ["Name", "Roll No.", "Class"] + [f"Subject {i} Marks" for i in range(1, SUBJECT_COUNT + 1)]
        self.add_vars = {}

        for i, field in enumerate(fields, start=1):
            ttk.Label(self.container, text=field + ":").grid(row=i, column=0, sticky="w", pady=5, padx=(0, 10))
            var = tk.StringVar()
            entry = ttk.Entry(self.container, textvariable=var, width=30)
            entry.grid(row=i, column=1, sticky="w", pady=5)
            self.add_vars[field] = var

        save_btn = ttk.Button(self.container, text="💾 Save", command=self._save_student)
        save_btn.grid(row=len(fields) + 1, column=0, columnspan=2, pady=18)

        self.add_result_label = ttk.Label(self.container, text="")
        self.add_result_label.grid(row=len(fields) + 2, column=0, columnspan=2, sticky="w")

    def _save_student(self):
        v = self.add_vars

        name = v["Name"].get().strip()
        roll_no = v["Roll No."].get().strip()
        student_class = v["Class"].get().strip()

        if not name or not roll_no or not student_class:
            messagebox.showwarning("Student Result", "Please fill Name, Roll No. and Class.")
            return

        if roll_no_exists(roll_no):
            messagebox.showwarning("Student Result", f"Roll No. {roll_no} already exists.")
            return

        marks = []
        for i in range(1, SUBJECT_COUNT + 1):
            raw = v[f"Subject {i} Marks"].get().strip()
            try:
                mark = float(raw)
            except ValueError:
                messagebox.showwarning("Student Result", f"Subject {i} marks must be a number.")
                return
            if mark < 0 or mark > MAX_MARKS_PER_SUBJECT:
                messagebox.showwarning("Student Result", f"Subject {i} marks must be between 0 and {MAX_MARKS_PER_SUBJECT}.")
                return
            marks.append(mark)

        total = sum(marks)
        max_total = SUBJECT_COUNT * MAX_MARKS_PER_SUBJECT
        percentage = round((total / max_total) * 100, 2)
        result = "Pass" if percentage >= PASS_PERCENTAGE else "Fail"

        row = [name, roll_no, student_class] + marks + [total, percentage, result]
        append_student_row(row)

        self.add_result_label.config(
            text=f"Saved — Total: {total}, Percentage: {percentage}%, Result: {result}",
            style="Pass.TLabel" if result == "Pass" else "Fail.TLabel",
        )

        for field in v:
            v[field].set("")

    # ==================================================================
    # 2. GET RESULT
    # ==================================================================
    def show_get_result(self):
        self._clear_container()

        ttk.Label(self.container, text="Get Result", style="Title.TLabel").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 14))

        ttk.Label(self.container, text="Enter Roll No.:").grid(row=1, column=0, sticky="w", padx=(0, 10))
        self.search_roll_var = tk.StringVar()
        entry = ttk.Entry(self.container, textvariable=self.search_roll_var, width=20)
        entry.grid(row=1, column=1, sticky="w")
        entry.bind("<Return>", lambda e: self._get_result())

        ttk.Button(self.container, text="🔍 Get Result", command=self._get_result).grid(row=1, column=2, padx=10)

        self.result_frame = ttk.Frame(self.container)
        self.result_frame.grid(row=2, column=0, columnspan=3, sticky="w", pady=20)

    def _get_result(self):
        for widget in self.result_frame.winfo_children():
            widget.destroy()

        roll_no = self.search_roll_var.get().strip()
        if not roll_no:
            messagebox.showwarning("Student Result", "Please enter a Roll No.")
            return

        record = find_student_by_roll(roll_no)
        if not record:
            ttk.Label(self.result_frame, text="❌ Student record not found.", style="Fail.TLabel").pack(anchor="w")
            return

        name, roll, student_class = record[0], record[1], record[2]
        total, percentage, result = record[-3], record[-2], record[-1]

        columns = ("name", "roll", "cls", "total", "pct", "result")
        tree = ttk.Treeview(self.result_frame, columns=columns, show="headings", height=1)
        headings = {"name": "Name", "roll": "Roll No.", "cls": "Class", "total": "Total Marks", "pct": "Percentage", "result": "Result"}
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=110, anchor="center")
        tree.pack()
        tree.insert("", "end", values=(name, roll, student_class, total, f"{percentage}%", result))

    # ==================================================================
    # 3. SHOW ALL RESULTS
    # ==================================================================
    def show_all_results(self):
        self._clear_container()

        ttk.Label(self.container, text="Show All Results", style="Title.TLabel").pack(anchor="w", pady=(0, 14))

        columns = ("name", "roll", "cls", "total", "pct", "result")
        tree = ttk.Treeview(self.container, columns=columns, show="headings")
        headings = {"name": "Name", "roll": "Roll No.", "cls": "Class", "total": "Total Marks", "pct": "Percentage", "result": "Result"}
        widths = {"name": 160, "roll": 90, "cls": 90, "total": 100, "pct": 100, "result": 90}
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], anchor="center" if col != "name" else "w")
        tree.pack(fill="both", expand=True)

        sb = ttk.Scrollbar(self.container, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")

        records = read_all_students()
        if not records:
            ttk.Label(self.container, text="No students saved yet.").pack(anchor="w", pady=10)
            return

        for record in records:
            name, roll, student_class = record[0], record[1], record[2]
            total, percentage, result = record[-3], record[-2], record[-1]
            tree.insert("", "end", values=(name, roll, student_class, total, f"{percentage}%", result))


if __name__ == "__main__":
    app = ResultApp()
    app.mainloop()